"""Canva-friendly PPTX pieces for the matching carousels (Mercedes's Canva test, 2026-10-19: the round photo came in
square, the doodles/burst/arrows were pictures she couldn't recolour, the built-in words were pictures).
Canva turns every PPTX shape into a path with a fill, so:
  photo_shape()  a photo slot as a shape (ellipse, arch or rectangle) filled with the picture: Canva keeps the
                 shape as the frame and Replace swaps the photo inside it
  vector_shape() her doodles traced to bezier paths (potrace) with a solid fill: recolourable in Canva
  lettering()    live text in DM Serif Display / Bodoni Moda (Canva fonts) instead of a picture; outlines are
                 added in Canva (they do not survive the import) and kept by the template link
"""
import copy, io
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P_ = 'http://schemas.openxmlformats.org/presentationml/2006/main'
R_ = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
EMU = 9525

def _e(tag, **attrs):
    ns, name = tag.split(':'); el = etree.Element('{%s}%s' % ({'a': A, 'p': P_, 'r': R_}[ns], name))
    for k, v in attrs.items():
        if ':' in k: k = '{%s}%s' % ({'r': R_}[k.split(':')[0]], k.split(':')[1])
        el.set(k, str(v))
    return el

def _sub(parent, tag, **attrs): el = _e(tag, **attrs); parent.append(el); return el

def _sp(slide, name, x, y, w, h, rot):
    """An empty p:sp at (x, y, w, h) px, rotated `rot` degrees clockwise; returns (sp, spPr)."""
    tree = slide.shapes._spTree; sid = max([int(e.get('id')) for e in tree.iter(qn('p:cNvPr'))] + [1]) + 1
    sp = _e('p:sp'); nv = _sub(sp, 'p:nvSpPr'); _sub(nv, 'p:cNvPr', id=sid, name=name); _sub(nv, 'p:cNvSpPr'); _sub(nv, 'p:nvPr')
    spPr = _sub(sp, 'p:spPr'); xf = _sub(spPr, 'a:xfrm', rot=int(round(rot * 60000)))
    _sub(xf, 'a:off', x=int(x * EMU), y=int(y * EMU)); _sub(xf, 'a:ext', cx=int(w * EMU), cy=int(h * EMU))
    tree.append(sp); return sp, spPr

def _geom(spPr, shape):
    if shape in ('circle', 'oval'): g = _sub(spPr, 'a:prstGeom', prst='ellipse'); _sub(g, 'a:avLst')
    elif shape == 'arch':
        g = _sub(spPr, 'a:prstGeom', prst='round2SameRect'); av = _sub(g, 'a:avLst'); _sub(av, 'a:gd', name='adj1', fmla='val 50000'); _sub(av, 'a:gd', name='adj2', fmla='val 0')
    else: g = _sub(spPr, 'a:prstGeom', prst='rect'); _sub(g, 'a:avLst')

def photo_shape(slide, img, name, c, size, rot=0, shape='rect'):
    """A photo slot: a shape whose fill is the picture (stretched to the shape, like a Canva frame)."""
    w, h = size; buf = io.BytesIO(); img.convert('RGB').save(buf, 'JPEG', quality=88); buf.seek(0)
    image_part, rId = slide.part.get_or_add_image_part(buf)
    sp, spPr = _sp(slide, name, c[0] - w / 2, c[1] - h / 2, w, h, rot); _geom(spPr, shape)
    bf = _sub(spPr, 'a:blipFill', rotWithShape='1'); _sub(bf, 'a:blip', **{'r:embed': rId}); st = _sub(bf, 'a:stretch'); _sub(st, 'a:fillRect')
    ln = _sub(spPr, 'a:ln'); _sub(ln, 'a:noFill')
    return sp

def trace(alpha, turdsize=12, scale=1.0):
    """Her doodle's ink (an L image, 255 = ink) as subpaths of cubic beziers in px: [[('M', p) | ('C', p1, p2, p) | ('L', p)]]."""
    import numpy as np, potrace
    bm = potrace.Bitmap(np.array(alpha) <= 128)    # potracer fills the False pixels (checked: > 128 traced the box around the ink)
    plist = bm.trace(turdsize=turdsize, opttolerance=0.4)
    paths = []
    for curve in plist:
        sp_ = [('M', (curve.start_point.x * scale, curve.start_point.y * scale))]
        for seg in curve.segments:
            if seg.is_corner: sp_ += [('L', (seg.c.x * scale, seg.c.y * scale)), ('L', (seg.end_point.x * scale, seg.end_point.y * scale))]
            else: sp_.append(('C', (seg.c1.x * scale, seg.c1.y * scale), (seg.c2.x * scale, seg.c2.y * scale), (seg.end_point.x * scale, seg.end_point.y * scale)))
        paths.append(sp_)
    return paths

def vector_shape(slide, alpha, name, c, rot, color, turdsize=12):
    """A doodle as one freeform shape (all its subpaths, holes by even-odd) with a solid, recolourable fill."""
    w, h = alpha.size; paths = trace(alpha, turdsize)
    sp, spPr = _sp(slide, name, c[0] - w / 2, c[1] - h / 2, w, h, rot)
    cg = _sub(spPr, 'a:custGeom'); [_sub(cg, t) for t in ('a:avLst', 'a:gdLst', 'a:ahLst', 'a:cxnLst')]
    _sub(cg, 'a:rect', l='0', t='0', r='r', b='b'); pl = _sub(cg, 'a:pathLst')
    k = EMU; p = _sub(pl, 'a:path', w=int(w * k), h=int(h * k))
    for sub in paths:
        for op in sub:
            if op[0] == 'M': m = _sub(p, 'a:moveTo'); _sub(m, 'a:pt', x=int(op[1][0] * k), y=int(op[1][1] * k))
            elif op[0] == 'L': m = _sub(p, 'a:lnTo'); _sub(m, 'a:pt', x=int(op[1][0] * k), y=int(op[1][1] * k))
            else:
                m = _sub(p, 'a:cubicBezTo')
                for q in op[1:]: _sub(m, 'a:pt', x=int(q[0] * k), y=int(q[1] * k))
        _sub(p, 'a:close')
    sf = _sub(spPr, 'a:solidFill'); _sub(sf, 'a:srgbClr', val='%02X%02X%02X' % tuple(color[:3]))
    ln = _sub(spPr, 'a:ln'); _sub(ln, 'a:noFill')
    return sp

FD = __import__('pathlib').Path(__file__).resolve().parent.parent / 'assets' / 'fonts'
LETTER_FONTS = {'dms': ('DM Serif Display', FD / 'dmserifdisplay-DMSerifDisplay-Regular.ttf', FD / 'dmserifdisplay-DMSerifDisplay-Italic.ttf'),
                'bodoni': ('Bodoni Moda', FD / 'BodoniModa.ttf', FD / 'BodoniModa.ttf')}

def lettering(slide, words, px, c, rot, fill, italic_last=True, font='dms'):
    """A built-in word as LIVE text in a Canva font (words, font and colour all editable). Canva drops text outlines on
    import (tested 2026-10-19), so the outline is added in Canva (Effects > Outline) before the template link is made:
    the list is written by carousels.py (canva-outline-list.md). Box one line tall (line height 1.2), centred."""
    from PIL import ImageFont
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    name, reg, ita = LETTER_FONTS[font]
    ws = words.split(' '); runs = [(w_ + (' ' if i < len(ws) - 1 else ''), bool(italic_last and i == len(ws) - 1 and len(ws) > 1)) for i, w_ in enumerate(ws)]
    fonts = {False: ImageFont.truetype(str(reg), px), True: ImageFont.truetype(str(ita), px)}
    adv = sum(fonts[it].getlength(t) for t, it in runs); w = int(adv + px * 0.4); h = int(px * 1.2)
    E = lambda v: Emu(int(round(v * EMU)))
    tb = slide.shapes.add_textbox(E(c[0] - w / 2), E(c[1] - h / 2), E(w), E(h)); tb.rotation = -rot; tb.name = 'LETTERING: ' + words
    tf = tb.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.TOP
    for mm in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, mm, 0)
    para = tf.paragraphs[0]; para.alignment = PP_ALIGN.CENTER; para.line_spacing = 1.2
    for t, it in runs:
        r = para.add_run(); r.text = t; f = r.font; f.name = name; f.size = Pt(px * 0.75); f.italic = it and font == 'dms'; f.color.rgb = RGBColor(*fill)
    return tb

def ink_alpha(rgba):
    """The ink of a recoloured doodle (its alpha), cropped tight: (L image, offset of the crop inside the original)."""
    a = rgba.split()[3]; bb = a.getbbox(); return a.crop(bb), bb
