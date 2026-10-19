"""A two-slide Canva import test for the carousel fixes (python build/canva_test.py -> build/.cache/canva-test.pptx):
slide 1 round, arched and tilted photo frames (shapes filled with a picture) + her doodles as recolourable vectors;
slide 2 the built-in words as live DM Serif Display text with an outline."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import carousel_kit as K
import canva_shapes as V
from collage_kit import doodle
from motion_b2 import Clip
import storyboard_b2 as S
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

E = lambda v: Emu(int(round(v * 9525)))

def serif_text(slide, words, px, c, rot, fill, edge, italic_last=False, w=1000):
    """Live DM Serif Display words with an outline (a:ln on the run) and no box."""
    h = int(px * 1.4); tb = slide.shapes.add_textbox(E(c[0] - w / 2), E(c[1] - h / 2), E(w), E(h)); tb.rotation = -rot
    tf = tb.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    para = tf.paragraphs[0]; para.alignment = PP_ALIGN.CENTER; ws = words.split(' ')
    for i, wd in enumerate(ws):
        r = para.add_run(); r.text = wd + (' ' if i < len(ws) - 1 else ''); f = r.font
        f.name = 'DM Serif Display'; f.size = Pt(px * 0.75); f.color.rgb = RGBColor(*fill); f.italic = bool(italic_last and i == len(ws) - 1)
        if edge:
            rPr = r._r.get_or_add_rPr(); ln = etree.SubElement(rPr, '{%s}ln' % V.A, w=str(int(max(3, px // 55) * 2 * 9525)))
            sf = etree.SubElement(ln, '{%s}solidFill' % V.A); etree.SubElement(sf, '{%s}srgbClr' % V.A, val='%02X%02X%02X' % edge)
            rPr.remove(ln); rPr.insert(0, ln)
    return tb

def main():
    prs = Presentation(); prs.slide_width, prs.slide_height = E(K.CW), E(K.CH); lay = prs.slide_layouts[6]
    paper = K.fit_bg(S.raw('physical', 'Blank Ruled Notebook Paper.png')); bgp = Path(__file__).parent / '.cache' / 'canva-test-bg.jpg'
    bgp.parent.mkdir(exist_ok=True); paper.convert('RGB').save(bgp, quality=88)
    prod = Clip('IMG_1381', 1.0, 6.0)
    s1 = prs.slides.add_slide(lay); s1.shapes.add_picture(str(bgp), 0, 0, E(K.CW), E(K.CH))
    V.photo_shape(s1, prod.at(2.0, (500, 500)), 'PHOTO: round', (300, 330), (480, 480), 0, 'circle')
    V.photo_shape(s1, prod.at(1.0, (400, 560)), 'PHOTO: arch', (800, 380), (380, 540), 0, 'arch')
    V.photo_shape(s1, prod.at(3.0, (560, 420)), 'PHOTO: tilted rectangle', (360, 900), (520, 400), -4, 'rect')
    for (kind, i, w, xy, rot) in (('burst', 3, 420, (820, 900), 0), ('arrow', 9, 220, (560, 640), 0), ('icon', 10, 120, (930, 1250), 0), ('icon', 4, 220, (300, 1250), 8), ('under', 0, 500, (600, 1330), -2)):
        a, _ = V.ink_alpha(doodle(kind, i, w, K.WINE)); V.vector_shape(s1, a, f'doodle {kind}', xy, rot, K.WINE)
    s2 = prs.slides.add_slide(lay); s2.shapes.add_picture(str(bgp), 0, 0, E(K.CW), E(K.CH))
    V.lettering(s2, 'this is for you', 150, (540, 280), 3, K.WHITE, K.WINE)
    V.lettering(s2, 'meet', 220, (540, 560), 0, K.PINK, K.WINE)
    V.lettering(s2, 'Comment', 200, (540, 860), 0, K.WINE, K.WHITE, edge_w=22)
    V.lettering(s2, 'made for', 150, (540, 1160), -2, K.WHITE, K.WINE)
    out = Path(__file__).parent / '.cache' / 'canva-test.pptx'; prs.save(out); print(out)

if __name__ == '__main__':
    main()
