"""The shared parts of the matching carousels (one per Reel template; recipes in carousels.py).

A carousel is 1080 x 1440 slides (3:4, the vault's carousel size), each a background (the Reel's stage) plus layers:
  art    built-in artwork (real objects with play, doodles, DM Serif lettering): a PNG the buyer leaves alone
  photo  a photo slot the buyer swaps (Replace in Canva); previews show stand-in stills from her July clips
  text   editable words in Montserrat Bold (in Canva's library and CapCut's), on a highlight box or bare
  guide  blue notes: "everything blue is a guide, delete it before you download"
Delivery: a .pptx the buyer drags into Canva Free (desktop browser), the path proven on the $1K Challenge library.
Carousel rules: automation/carousels/content-rules.md, minus the Scrolly brand colours (FRAMEWORK "Matching carousels").
"""
import math, sys
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont, ImageOps
sys.path.insert(0, str(Path(__file__).resolve().parent))
import collage as C
import play as P
import storyboard_b2 as S
import story7 as S7
from lettering_b2 import FONT

CW, CH = 1080, 1440
WINE, PINK, BUTTER, BLUE = P.PAL['wine'], P.PAL['pink'], P.PAL['butter'], P.PAL['blue']
WHITE, INK = (255, 255, 255), (24, 23, 22)
GUIDE = (37, 99, 235)
MONT = FONT['montserrat']
CAPTION = ('Caption to start from (CTA in the first two sentences, 3 hashtags): Comment [KEYWORD] and I will send you '
           '[the link]. [One line of real context: who it helps or how you use it.] #[tag] #[tag] #[tag]')
FIRST_GUIDE = 'Everything blue is a guide. Delete it before you download.'

# ---- layers ------------------------------------------------------------------------------------------------------------
def art(img, c, rot=0, name='art'): return dict(kind='art', img=img, c=c, rot=rot, name=name)
def photo(slot, c, size, rot=0, still=None, shape='rect', gray=False):
    return dict(kind='photo', slot=slot, c=c, size=(int(size[0]), int(size[1])), rot=rot, still=still, shape=shape, gray=gray)
def text(words, px, c, rot=0, box=WHITE, fg=INK, pill=False, align='center'):
    return dict(kind='text', words=words, px=px, c=c, rot=rot, box=box, fg=fg, pill=pill, align=align)
def guide(words, c=(540, 1390), px=26): return dict(kind='guide', words=words, px=px, c=c, rot=0, box=GUIDE, fg=WHITE, pill=False, align='center')

def mfont(px): return ImageFont.truetype(MONT, px)

def text_box(L):
    """The words' box in px, the same numbers the .pptx uses: (w, h, pad_x, pad_y, line_h)."""
    f = mfont(L['px']); lines = L['words'].split('\n'); lh = L['px'] * 1.22
    pad_x = L['px'] * (0.55 if L['pill'] else 0.28); pad_y = L['px'] * (0.22 if L['pill'] else 0.12)
    w = max(f.getlength(l) for l in lines) + 2 * pad_x; h = lh * len(lines) + 2 * pad_y
    return int(w), int(h), pad_x, pad_y, lh

def text_img(L):
    w, h, px_, py_, lh = text_box(L); f = mfont(L['px']); im = Image.new('RGBA', (w, h), (0, 0, 0, 0)); d = ImageDraw.Draw(im)
    if L['box']:
        if L['pill']: d.rounded_rectangle((0, 0, w - 1, h - 1), h // 2, fill=L['box'] + (255,))
        else: d.rectangle((0, 0, w - 1, h - 1), fill=L['box'] + (255,))
    x, an = {'center': (w / 2, 'mm'), 'left': (px_, 'lm'), 'right': (w - px_, 'rm')}[L['align']]
    for k, l in enumerate(L['words'].split('\n')): d.text((x, py_ + lh * (k + 0.5)), l, font=f, fill=L['fg'], anchor=an)
    return im

# ---- artwork helpers ---------------------------------------------------------------------------------------------------
def word(t, px, color=WHITE, edge=None, italic_last=True):
    """A built-in DM Serif line (last word Italic), words set tight like the Reels' lettering."""
    ws = t.split(); ims = [P.story_word(w_, px, color, edge=edge or WINE, italic=(italic_last and i == len(ws) - 1 and len(ws) > 1)) for i, w_ in enumerate(ws)]
    step = [im.width - 80 + int(px * 0.3) for im in ims]; out = Image.new('RGBA', (sum(step[:-1]) + ims[-1].width, max(im.height for im in ims)), (0, 0, 0, 0)); x = 0
    for im, st in zip(ims, step): out.alpha_composite(im, (x, (out.height - im.height) // 2)); x += st
    return out

def fit_bg(im): return ImageOps.fit(im, (CW, CH), Image.LANCZOS).convert('RGBA')

def at(c, rot, local):
    """A point given as an offset from an art layer's centre (before rotation), in slide coordinates."""
    a = math.radians(rot); dx, dy = local
    return (c[0] + dx * math.cos(a) + dy * math.sin(a), c[1] - dx * math.sin(a) + dy * math.cos(a))

def torn_card(size, seed, margin=26, color=(252, 250, 245)):
    """Torn paper behind a photo slot: its ragged edge shows round the photo."""
    w, h = size[0] + 2 * margin, size[1] + 2 * margin; card = Image.new('RGBA', (w, h), color + (255,))
    card.putalpha(C.torn_mask((w, h), seed, 14, 'tblr')); return C.shadow(card, (8, 14), 14, 100)

def white_card(size, border=18):
    """A white photo-print border behind a photo slot (the photo sits on top, inset by `border`)."""
    c = Image.new('RGBA', (size[0] + 2 * border, size[1] + 2 * border), WHITE + (255,)); return C.shadow(c, (8, 14), 14, 110)

def round_card(d, margin=18):
    im = Image.new('RGBA', (d + 2 * margin, d + 2 * margin), (0, 0, 0, 0)); ImageDraw.Draw(im).ellipse((0, 0, d + 2 * margin - 1, d + 2 * margin - 1), fill=(252, 250, 245, 255))
    return C.shadow(im, (8, 14), 14, 100)

def torn_strip(bg, h, seed, edge='b'):
    """A strip of a stage torn along one edge ('b': a top strip torn at its bottom; 't': a bottom strip torn at its top)."""
    strip = bg.crop((0, 0, CW, h) if edge == 'b' else (0, CH - h, CW, CH)); strip.putalpha(C.torn_mask((CW, h), seed, 16, edge))
    return C.shadow(strip, (0, 12 if edge == 'b' else -12), 14, 110)

def string(a, b, sag=40, color=(120, 108, 100), width=4):
    """A string from a to b (slide coordinates) as its own art layer: (image, centre)."""
    x0, y0 = min(a[0], b[0]) - abs(sag) - 10, min(a[1], b[1]) - 10; x1, y1 = max(a[0], b[0]) + abs(sag) + 10, max(a[1], b[1]) + 10
    im = Image.new('RGBA', (int(x1 - x0), int(y1 - y0)), (0, 0, 0, 0)); mid = ((a[0] + b[0]) / 2 + sag, (a[1] + b[1]) / 2)
    pts = [((1 - u) ** 2 * a[0] + 2 * (1 - u) * u * mid[0] + u * u * b[0] - x0, (1 - u) ** 2 * a[1] + 2 * (1 - u) * u * mid[1] + u * u * b[1] - y0) for u in [i / 24 for i in range(25)]]
    ImageDraw.Draw(im).line(pts, fill=color + (255,), width=width)
    return im, ((x0 + x1) / 2, (y0 + y1) / 2)

def device_art(w, name=None):
    """A real green-screen device with a dark screen: (art, screen box inside the art) for a photo slot on its screen."""
    im, m, bb = P.device(name or S7.PHONE, w); im = im.copy(); im.paste(Image.new('RGBA', im.size, (20, 20, 22, 255)), (0, 0), m)
    px = im.load()
    for y in range(im.height):
        for x in range(im.width):
            r, g_, b_, a = px[x, y]
            if a and g_ > 120 and g_ > r + 40 and g_ > b_ + 40: px[x, y] = (20, 20, 22, a)
    edge = 8; cut = P.die_cut(im, edge, shadow=False); shadowed = C.shadow(cut, (8, 14), 14, 90)
    off = (edge + 4) + (shadowed.width - cut.width) // 2
    return shadowed, (bb[0] + off, bb[1] + off, bb[2] + off, bb[3] + off)

def on_screen(c, art_img, box, still, slot, inset=4):
    """The photo slot that fills a device's (or frame's) window, for art centred at c (unrotated)."""
    x0, y0, x1, y1 = box; cx = c[0] - art_img.width / 2 + (x0 + x1) / 2; cy = c[1] - art_img.height / 2 + (y0 + y1) / 2
    return photo(slot, (cx, cy), (x1 - x0 - 2 * inset, y1 - y0 - 2 * inset), 0, still)

def blank_card(lines, w, accent, kind='dm', pin='pin-head-black', name='[customer name hidden]'):
    """A pinned DM/review card with an empty bubble (its words are an editable layer): (art, bubble centre offset, px)."""
    card = P.proof_card([' '] * lines, w, accent, name, kind); pn = int(w * 0.13); pinned = P.pinned(card, 0, pin)
    pad = (pinned.width - card.width) // 2; px = int(w * 0.065); lh = int(px * 1.3); bh = lh * lines + int(w * 0.12); by = int(w * 0.24)
    cy = pad + pn // 2 + by + bh / 2; return pinned, (0, cy - pinned.height / 2), px

def note_text(c, rot, words, px=50, k=(0, 22), fg=INK, align='center'):
    """Editable words on a sticky note art centred at c (its paper sits a little below the art's centre)."""
    return text(words, px, at(c, rot, k), rot, box=None, fg=fg, align=align)

# ---- preview PNGs ----------------------------------------------------------------------------------------------------
def place(im, item, c, rot):
    it = item.rotate(rot, expand=True, resample=Image.BICUBIC) if rot else item
    im.alpha_composite(it, (int(c[0] - it.width / 2), int(c[1] - it.height / 2)))

def shape_mask(shape, size):
    w, h = size; m = Image.new('L', size, 0); d = ImageDraw.Draw(m)
    if shape in ('circle', 'oval'): d.ellipse((0, 0, w - 1, h - 1), fill=255)
    elif shape == 'arch': d.rectangle((0, w // 2, w, h), fill=255); d.pieslice((0, 0, w - 1, w - 1), 180, 360, fill=255)
    else: d.rectangle((0, 0, w, h), fill=255)
    return m

def photo_img(L, still):
    v = ImageOps.fit(still.convert('RGBA'), L['size'])
    if L['gray']: v = ImageOps.grayscale(v).convert('RGBA')
    if L['shape'] != 'rect': v.putalpha(shape_mask(L['shape'], v.size))
    return v

def placeholder(L):
    w, h = L['size']; im = Image.new('RGB', (w, h), (236, 226, 230)); d = ImageDraw.Draw(im)
    d.text((w / 2, h / 2), 'YOUR PHOTO', font=mfont(max(22, min(w // 12, 60))), fill=WINE, anchor='mm')
    return im

def render(sl, guides=False, stills=True):
    im = sl['bg'].copy()
    for L in sl['layers']:
        if L['kind'] == 'art': place(im, L['img'], L['c'], L['rot'])
        elif L['kind'] == 'photo': place(im, photo_img(L, L['still'] if stills else placeholder(L)), L['c'], L['rot'])
        elif L['kind'] == 'text' or (guides and L['kind'] == 'guide'): place(im, text_img(L), L['c'], L['rot'])
    return im.convert('RGB')

# ---- the .pptx ---------------------------------------------------------------------------------------------------------
def pptx(ss, path, tmp):
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    E = lambda v: Emu(int(round(v * 9525)))
    prs = Presentation(); prs.slide_width, prs.slide_height = E(CW), E(CH); blank = prs.slide_layouts[6]
    tmp.mkdir(parents=True, exist_ok=True); n = 0; bgs = {}
    for i, sl in enumerate(ss, 1):
        slide = prs.slides.add_slide(blank)
        if id(sl['bg']) not in bgs:
            bgs[id(sl['bg'])] = tmp / f'bg{len(bgs)}.jpg'; sl['bg'].convert('RGB').save(bgs[id(sl['bg'])], quality=90)
        slide.shapes.add_picture(str(bgs[id(sl['bg'])]), 0, 0, E(CW), E(CH)).name = 'background'
        for L in sl['layers']:
            if L['kind'] in ('art', 'photo'):
                n += 1; p = tmp / f'{n:03d}.png'
                im = L['img'] if L['kind'] == 'art' else placeholder(L)
                im.save(p); w, h = (im.width, im.height) if L['kind'] == 'art' else L['size']
                pic = slide.shapes.add_picture(str(p), E(L['c'][0] - w / 2), E(L['c'][1] - h / 2), E(w), E(h))
                pic.rotation = -L['rot']; pic.name = L['name'] if L['kind'] == 'art' else 'PHOTO: ' + L['slot']
                if L['kind'] == 'photo' and L['shape'] in ('circle', 'oval'): pic.auto_shape_type = MSO_SHAPE.OVAL
                elif L['kind'] == 'photo' and L['shape'] == 'arch':
                    pic.auto_shape_type = MSO_SHAPE.ROUND_2_SAME_RECTANGLE; geom = pic._element.spPr.prstGeom   # rounded top = an arch
                    av = geom.find('{http://schemas.openxmlformats.org/drawingml/2006/main}avLst')
                    if av is None: av = geom.makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}avLst', {}); geom.append(av)
                    av.append(av.makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}gd', {'name': 'adj1', 'fmla': 'val 50000'}))
            else:
                w, h, pad_x, pad_y, lh = text_box(L)
                shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if L['pill'] else MSO_SHAPE.RECTANGLE,
                                             E(L['c'][0] - w / 2), E(L['c'][1] - h / 2), E(w), E(h))
                if L['pill']: shp.adjustments[0] = 0.5
                shp.rotation = -L['rot']; shp.line.fill.background(); shp.shadow.inherit = False
                if L['box']: shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor(*L['box'])
                else: shp.fill.background()
                tf = shp.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf.margin_left = tf.margin_right = E(pad_x); tf.margin_top = tf.margin_bottom = E(pad_y)
                for k, line in enumerate(L['words'].split('\n')):
                    para = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); para.line_spacing = Pt(lh * 0.75)
                    para.alignment = {'center': PP_ALIGN.CENTER, 'left': PP_ALIGN.LEFT, 'right': PP_ALIGN.RIGHT}[L['align']]
                    r = para.add_run(); r.text = line; r.font.name = 'Montserrat'; r.font.bold = True; r.font.size = Pt(L['px'] * 0.75); r.font.color.rgb = RGBColor(*L['fg'])
                shp.name = 'GUIDE (delete)' if L['kind'] == 'guide' else 'TEXT: ' + L['words'].replace('\n', ' ')
        slide.notes_slide.notes_text_frame.text = f'Slide {i}. ' + sl['note']
    prs.save(str(path))

def sheet(ims, path):
    th = [im.resize((360, 480), Image.LANCZOS) for im in ims]; out = Image.new('RGB', (len(th) * 372 + 12, 492), (255, 255, 255))
    for i, t in enumerate(th): out.paste(t, (12 + i * 372, 6))
    out.save(path, quality=88)

def build(ss, out, name, cache):
    """Slide PNGs, a contact sheet, the as-the-buyer-opens-it sheet and the .pptx into `out`."""
    out.mkdir(parents=True, exist_ok=True); (out / 'slides').mkdir(exist_ok=True)
    for old in (out / 'slides').glob('*.png'): old.unlink()
    ims = [render(sl) for sl in ss]
    for i, im in enumerate(ims, 1): im.save(out / 'slides' / f'{i:02d}.png')
    sheet(ims, out / 'contact-sheet.jpg'); sheet([render(sl, True, False) for sl in ss], out / 'as-the-buyer-opens-it.jpg')
    pptx(ss, out / f'{name}.pptx', cache)
